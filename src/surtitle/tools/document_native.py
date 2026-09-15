"""Read and write documents with nothing installed but the application itself.

The point of this module is that document handling works on a machine that has
never heard of LibreOffice. It reads the formats people actually exchange
(Word, Excel, PowerPoint, OpenDocument, PDF, and the plain ones) and writes the
formats they actually ask for (PDF, text, CSV, HTML, XLSX).

The design is deliberately two halves. An *extractor* turns any supported source
into a neutral :class:`Document` — a flat list of headings, paragraphs, list
items, code and tables — and a *renderer* writes that document out in one target
format. Supporting a new format therefore means writing one function on one side
rather than a converter for every source/target pair.

Fidelity is honest rather than high: this reads *content*, not layout. Fonts,
colours, columns, headers and footers, embedded images and tracked changes are
dropped. LibreOffice keeps them, which is why
:mod:`surtitle.tools.documents` still offers it when it happens to be installed
— it is a quality upgrade, not a requirement.
"""

from __future__ import annotations

import csv
import json
import re
import zipfile
from collections.abc import Callable, Iterator
from dataclasses import dataclass, field
from html.parser import HTMLParser
from pathlib import Path
from typing import Any, ClassVar
from xml.etree import ElementTree

__all__ = [
    "OFFICE_SUFFIXES",
    "SOURCE_SUFFIXES",
    "TARGETS",
    "Block",
    "Document",
    "NativeConversionError",
    "can_convert",
    "convert",
    "document_to_text",
    "extract",
]


class NativeConversionError(RuntimeError):
    """Raised when the built-in converter cannot handle a file."""


@dataclass(slots=True)
class Block:
    """One piece of a document.

    ``kind`` is one of ``heading``, ``paragraph``, ``list_item``, ``code`` or
    ``table``. Only the field that kind needs is populated; a table carries
    ``rows``, everything else carries ``text``.
    """

    kind: str
    text: str = ""
    level: int = 1
    rows: list[list[str]] = field(default_factory=list)


@dataclass(slots=True)
class Document:
    blocks: list[Block] = field(default_factory=list)
    title: str = ""


# --------------------------------------------------------------------------- #
# Which formats are understood, and how far to go with them.
# --------------------------------------------------------------------------- #

_PLAIN_SUFFIXES = frozenset({".txt", ".text", ".log", ".rst"})
_MARKDOWN_SUFFIXES = frozenset({".md", ".markdown"})
_HTML_SUFFIXES = frozenset({".html", ".htm", ".xhtml"})
_OOXML_SUFFIXES = frozenset({".docx", ".docm", ".pptx", ".xlsx", ".xlsm"})
_ODF_SUFFIXES = frozenset({".odt", ".ods", ".odp"})
_DELIMITED_SUFFIXES = {".csv": ",", ".tsv": "\t"}

# Zip-based containers. These are the formats a plain text read cannot handle,
# so `read_file` routes them through :func:`extract` instead of rejecting them
# as binary.
OFFICE_SUFFIXES = _OOXML_SUFFIXES | _ODF_SUFFIXES

SOURCE_SUFFIXES = (
    _PLAIN_SUFFIXES
    | _MARKDOWN_SUFFIXES
    | _HTML_SUFFIXES
    | _OOXML_SUFFIXES
    | _ODF_SUFFIXES
    | frozenset({".pdf", ".json"})
    | frozenset(_DELIMITED_SUFFIXES)
)

# Targets this module can write. Anything else is LibreOffice's business.
TARGETS = frozenset({"pdf", "txt", "csv", "html", "xlsx"})

# A spreadsheet can legitimately be enormous, and rendering one is not the same
# as reading it. These bounds keep an accidental "convert this 400k-row export"
# from turning into a multi-gigabyte PDF. Anything dropped is reported in the
# output rather than silently ignored.
_MAX_TABLE_ROWS = 5000
_MAX_TABLE_COLUMNS = 64

# A package is decompressed into memory, so a small file can still claim to be an
# enormous one. Refusing an absurd member keeps a malicious or corrupt document
# from exhausting memory, and the limit is far above any real document part.
_MAX_MEMBER_BYTES = 64 * 1024 * 1024

# A deck with thousands of slides is not a document anyone reads, and each slide
# adds blocks that are all held at once.
_MAX_SLIDES = 1000

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
_ODF_TEXT = "{urn:oasis:names:tc:opendocument:xmlns:text:1.0}"
_ODF_TABLE = "{urn:oasis:names:tc:opendocument:xmlns:table:1.0}"
_ODF_OFFICE = "{urn:oasis:names:tc:opendocument:xmlns:office:1.0}"
_DC = "{http://purl.org/dc/elements/1.1/}"


def can_convert(suffix: str, target: str) -> bool:
    """Whether the built-in converter handles this source/target pair."""
    return suffix.lower() in SOURCE_SUFFIXES and target.lower() in TARGETS


# --------------------------------------------------------------------------- #
# Small text helpers shared by every extractor.
# --------------------------------------------------------------------------- #


def _tidy(text: str) -> str:
    """Collapse the runs of whitespace that are an artefact of XML pretty-printing."""
    return re.sub(r"[ \t\r\f\v]+", " ", text).strip()


def _paragraphs(text: str) -> Iterator[str]:
    """Split extracted text into paragraphs on blank lines."""
    for chunk in re.split(r"\n\s*\n", text):
        cleaned = "\n".join(line.rstrip() for line in chunk.splitlines()).strip()
        if cleaned:
            yield cleaned


def _clip_columns(row: list[Any]) -> list[str]:
    """Normalise one row to strings, bounded to a sane number of columns."""
    cells = ["" if cell is None else _tidy(str(cell)) for cell in row[:_MAX_TABLE_COLUMNS]]
    if len(row) > _MAX_TABLE_COLUMNS:
        cells.append(f"... {len(row) - _MAX_TABLE_COLUMNS} more")
    return cells


def _clip_rows(rows: list[list[Any]]) -> list[list[str]]:
    """Bound a table, leaving a visible note where content was dropped."""
    clipped = [_clip_columns(row) for row in rows[:_MAX_TABLE_ROWS]]
    if len(rows) > _MAX_TABLE_ROWS:
        clipped.append([f"... {len(rows) - _MAX_TABLE_ROWS} more row(s) not shown"])
    return clipped


# --------------------------------------------------------------------------- #
# Extractors: source file -> Document
# --------------------------------------------------------------------------- #


def extract(path: Path) -> Document:
    """Read any supported source into the neutral model.

    Raises :class:`NativeConversionError` for a format this module does not know,
    or when the file cannot be parsed — the caller decides whether to fall back
    to LibreOffice.
    """
    suffix = path.suffix.lower()
    handler = _EXTRACTORS.get(suffix)
    if handler is None:
        raise NativeConversionError(
            f"{suffix or 'that file type'} is not readable by the built-in converter"
        )
    try:
        return handler(path)
    except NativeConversionError:
        raise
    except Exception as exc:
        raise NativeConversionError(
            f"could not read {path.name} as {suffix.lstrip('.')}: {type(exc).__name__}: {exc}"
        ) from exc


def _extract_plain(path: Path, *, markdown: bool = False) -> Document:
    text = path.read_text(encoding="utf-8", errors="replace")
    blocks: list[Block] = []
    for paragraph in _paragraphs(text):
        if markdown:
            heading = re.match(r"^(#{1,6})\s+(.*)$", paragraph)
            if heading and "\n" not in paragraph:
                level = len(heading.group(1))
                blocks.append(Block("heading", _tidy(heading.group(2)), level=level))
                continue
            if all(re.match(r"^\s*[-*+]\s+", line) for line in paragraph.splitlines()):
                for line in paragraph.splitlines():
                    blocks.append(Block("list_item", _tidy(re.sub(r"^\s*[-*+]\s+", "", line))))
                continue
        blocks.append(Block("paragraph", paragraph))
    return Document(blocks, title=path.stem)


def _extract_json(path: Path) -> Document:
    raw = path.read_text(encoding="utf-8", errors="replace")
    try:
        # Re-serialising gives the model a consistent shape to read even when the
        # file on disk was minified onto a single line.
        pretty = json.dumps(json.loads(raw), indent=2, ensure_ascii=False)
    except ValueError:
        pretty = raw
    return Document([Block("code", pretty)], title=path.stem)


def _extract_delimited(path: Path) -> Document:
    delimiter = _DELIMITED_SUFFIXES[path.suffix.lower()]
    rows: list[list[str]] = []
    with path.open("r", encoding="utf-8", errors="replace", newline="") as handle:
        for row in csv.reader(handle, delimiter=delimiter):
            rows.append(row)
    if not rows:
        return Document([], title=path.stem)
    return Document([Block("table", rows=_clip_rows(rows))], title=path.stem)


def _extract_pdf(path: Path) -> Document:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    blocks: list[Block] = []
    for page in reader.pages:
        for paragraph in _paragraphs(page.extract_text() or ""):
            blocks.append(Block("paragraph", paragraph))
    title = ""
    if reader.metadata is not None:
        title = _tidy(str(reader.metadata.title or ""))
    return Document(blocks, title=title or path.stem)


def _read_member(archive: zipfile.ZipFile, name: str) -> bytes:
    """Read one part of a package, refusing an implausibly large one."""
    info = archive.getinfo(name)
    if info.file_size > _MAX_MEMBER_BYTES:
        raise NativeConversionError(
            f"{name} expands to {info.file_size // (1024 * 1024)} MB, which is more "
            "than this converter will read from one document part"
        )
    return archive.read(name)


def _core_title(archive: zipfile.ZipFile) -> str:
    """Read the document title out of an OOXML package, if it declares one."""
    try:
        root = ElementTree.fromstring(_read_member(archive, "docProps/core.xml"))
    except (KeyError, ElementTree.ParseError):
        return ""
    node = root.find(f"{_DC}title")
    return _tidy(node.text or "") if node is not None else ""


def _ooxml_text(element: ElementTree.Element) -> str:
    """Concatenate the runs of text inside one OOXML paragraph."""
    parts: list[str] = []
    for node in element.iter():
        if node.tag == f"{_W}t":
            parts.append(node.text or "")
        elif node.tag == f"{_W}tab":
            parts.append("\t")
        elif node.tag in (f"{_W}br", f"{_W}cr"):
            parts.append("\n")
    return _tidy("".join(parts))


def _docx_block(element: ElementTree.Element) -> Block | None:
    text = _ooxml_text(element)
    if not text:
        return None
    style = element.find(f"{_W}pPr/{_W}pStyle")
    name = (style.get(f"{_W}val") or "") if style is not None else ""
    heading = re.fullmatch(r"[Hh]eading\s*(\d+)", name)
    if heading:
        return Block("heading", text, level=min(int(heading.group(1)), 6))
    if name.lower() in {"title", "subtitle"}:
        return Block("heading", text, level=1 if name.lower() == "title" else 2)
    if name.lower().startswith("list"):
        return Block("list_item", text)
    return Block("paragraph", text)


def _docx_table(element: ElementTree.Element) -> Block:
    rows: list[list[str]] = []
    for row in element.findall(f"{_W}tr"):
        cells = [_ooxml_text(cell) for cell in row.findall(f"{_W}tc")]
        if any(cells):
            rows.append(cells)
    return Block("table", rows=_clip_rows(rows))


def _extract_docx(path: Path) -> Document:
    with zipfile.ZipFile(path) as archive:
        root = ElementTree.fromstring(_read_member(archive, "word/document.xml"))
        title = _core_title(archive)
    body = root.find(f"{_W}body")
    blocks: list[Block] = []
    for child in body if body is not None else root:
        if child.tag == f"{_W}p":
            block = _docx_block(child)
            if block is not None:
                blocks.append(block)
        elif child.tag == f"{_W}tbl":
            table = _docx_table(child)
            if table.rows:
                blocks.append(table)
    return Document(blocks, title=title or path.stem)


def _extract_xlsx(path: Path) -> Document:
    from openpyxl import load_workbook

    # read_only streams the sheet rather than building the whole object graph,
    # which matters for the large exports this is most likely to be handed.
    book = load_workbook(filename=str(path), read_only=True, data_only=True)
    blocks: list[Block] = []
    try:
        sheets = list(book.worksheets)
        for sheet in sheets:
            rows: list[list[str]] = []
            truncated = False
            for index, row in enumerate(sheet.iter_rows(values_only=True)):
                if index >= _MAX_TABLE_ROWS:
                    truncated = True
                    break
                cells = ["" if value is None else str(value) for value in row]
                if any(cell.strip() for cell in cells):
                    rows.append(cells)
            if not rows:
                continue
            if len(sheets) > 1:
                blocks.append(Block("heading", sheet.title, level=2))
            clipped = [_clip_columns(row) for row in rows]
            if truncated:
                clipped.append([f"... more row(s) not shown (stopped at {_MAX_TABLE_ROWS})"])
            blocks.append(Block("table", rows=clipped))
    finally:
        book.close()
    return Document(blocks, title=path.stem)


def _slide_number(name: str) -> int:
    match = re.search(r"(\d+)\.xml$", name)
    return int(match.group(1)) if match else 0


def _extract_pptx(path: Path) -> Document:
    blocks: list[Block] = []
    with zipfile.ZipFile(path) as archive:
        title = _core_title(archive)
        slides = [n for n in archive.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)]
        ordered = sorted(slides, key=_slide_number)
        for index, name in enumerate(ordered[:_MAX_SLIDES], start=1):
            root = ElementTree.fromstring(_read_member(archive, name))
            texts = [t for t in (_tidy(node.text or "") for node in root.iter(f"{_A}t")) if t]
            if not texts:
                continue
            # A slide's first text frame is its title in every deck that has one.
            # Where it is not, the cost is one heading that reads like a sentence.
            blocks.append(Block("heading", texts[0], level=2))
            blocks.extend(Block("paragraph", text) for text in texts[1:])
            if index < len(ordered):
                blocks.append(Block("paragraph", ""))
        if len(ordered) > _MAX_SLIDES:
            blocks.append(Block("paragraph", f"... {len(ordered) - _MAX_SLIDES} more slide(s)"))
    return Document(blocks, title=title or path.stem)


def _odf_text(element: ElementTree.Element) -> str:
    return _tidy("".join(element.itertext()))


def _odf_walk(element: ElementTree.Element, blocks: list[Block]) -> None:
    """Walk ODF body content, treating a whole table as one block.

    Cells contain ``text:p`` elements of their own, so descending into a table
    would emit every cell twice — once as a table and once as loose paragraphs.
    """
    for child in element:
        if child.tag == f"{_ODF_TEXT}h":
            text = _odf_text(child)
            if text:
                raw_level = child.get(f"{_ODF_TEXT}outline-level") or "1"
                level = min(int(raw_level), 6) if raw_level.isdigit() else 1
                blocks.append(Block("heading", text, level=level))
        elif child.tag == f"{_ODF_TEXT}p":
            text = _odf_text(child)
            if text:
                blocks.append(Block("paragraph", text))
        elif child.tag == f"{_ODF_TABLE}table":
            rows: list[list[str]] = []
            for row in child.findall(f"{_ODF_TABLE}table-row"):
                cells = [_odf_text(cell) for cell in row.findall(f"{_ODF_TABLE}table-cell")]
                if any(cells):
                    rows.append(cells)
            if rows:
                blocks.append(Block("table", rows=_clip_rows(rows)))
        else:
            _odf_walk(child, blocks)


def _extract_odf(path: Path) -> Document:
    with zipfile.ZipFile(path) as archive:
        root = ElementTree.fromstring(_read_member(archive, "content.xml"))
    blocks: list[Block] = []
    body = root.find(f"{_ODF_OFFICE}body")
    _odf_walk(body if body is not None else root, blocks)
    return Document(blocks, title=path.stem)


class _HtmlReader(HTMLParser):
    """Turn HTML into blocks, keeping headings, lists and tables.

    Text inside a table is collected per cell rather than as loose paragraphs,
    which is why block-level tags are ignored while a table is open.
    """

    _HEADINGS: ClassVar[dict[str, int]] = {"h1": 1, "h2": 2, "h3": 3, "h4": 4, "h5": 5, "h6": 6}
    _BLOCKS: ClassVar[set[str]] = {"p", "li", "pre", "blockquote", "dd", "dt"}

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.blocks: list[Block] = []
        self._skip = 0
        self._table_depth = 0
        self._rows: list[list[str]] = []
        self._row: list[str] | None = None
        self._cell: list[str] | None = None
        self._kind: str | None = None
        self._level = 1
        self._text: list[str] = []

    # -- opening tags ------------------------------------------------------ #
    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in {"script", "style"}:
            self._skip += 1
            return
        if tag == "table":
            if self._table_depth == 0:
                self._rows = []
            self._table_depth += 1
            return
        if self._table_depth:
            if tag == "tr":
                self._row = []
            elif tag in {"td", "th"}:
                self._cell = []
            return
        if tag in self._HEADINGS:
            self._start("heading", self._HEADINGS[tag])
        elif tag in self._BLOCKS:
            self._start("code" if tag == "pre" else ("list_item" if tag == "li" else "paragraph"))
        elif tag == "br" and self._kind is not None:
            self._text.append("\n")

    def _start(self, kind: str, level: int = 1) -> None:
        if self._kind is not None:
            self._flush()
        self._kind, self._level, self._text = kind, level, []

    def _flush(self) -> None:
        text = _tidy("".join(self._text))
        if text and self._kind is not None:
            self.blocks.append(Block(self._kind, text, level=self._level))
        self._kind, self._text = None, []

    # -- closing tags ------------------------------------------------------ #
    def handle_endtag(self, tag: str) -> None:
        if tag in {"script", "style"}:
            self._skip = max(0, self._skip - 1)
            return
        if tag == "table":
            self._table_depth = max(0, self._table_depth - 1)
            if self._table_depth == 0 and any(any(c for c in row) for row in self._rows):
                self.blocks.append(Block("table", rows=_clip_rows(self._rows)))
            return
        if self._table_depth:
            if tag in {"td", "th"}:
                if self._row is not None:
                    self._row.append(_tidy("".join(self._cell or [])))
                self._cell = None
            elif tag == "tr":
                if self._row is not None and any(c for c in self._row):
                    self._rows.append(self._row)
                self._row = None
            return
        if tag in self._HEADINGS or tag in self._BLOCKS:
            self._flush()

    def handle_data(self, data: str) -> None:
        if self._skip:
            return
        if self._cell is not None:
            self._cell.append(data)
        elif self._kind is not None:
            self._text.append(data)

    def close(self) -> None:
        super().close()
        if not self._table_depth:
            self._flush()


def _extract_html(path: Path) -> Document:
    reader = _HtmlReader()
    reader.feed(path.read_text(encoding="utf-8", errors="replace"))
    reader.close()
    return Document(reader.blocks, title=path.stem)


_EXTRACTORS: dict[str, Callable[[Path], Document]] = {
    **dict.fromkeys(_PLAIN_SUFFIXES, _extract_plain),
    **{suffix: (lambda p: _extract_plain(p, markdown=True)) for suffix in _MARKDOWN_SUFFIXES},
    **dict.fromkeys(_DELIMITED_SUFFIXES, _extract_delimited),
    **dict.fromkeys(_HTML_SUFFIXES, _extract_html),
    ".json": _extract_json,
    ".pdf": _extract_pdf,
    ".docx": _extract_docx,
    ".docm": _extract_docx,
    ".xlsx": _extract_xlsx,
    ".xlsm": _extract_xlsx,
    ".pptx": _extract_pptx,
    **dict.fromkeys(_ODF_SUFFIXES, _extract_odf),
}


# --------------------------------------------------------------------------- #
# Renderers: Document -> target file
# --------------------------------------------------------------------------- #


def document_to_text(document: Document) -> str:
    """Flatten a document to readable plain text.

    This is also what `read_file` uses to make an Office document readable
    without converting it first.
    """
    lines: list[str] = []
    for block in document.blocks:
        if block.kind == "table":
            lines.extend("\t".join(row) for row in block.rows)
            lines.append("")
        elif block.kind == "list_item":
            lines.append(f"- {block.text}")
        elif block.kind == "heading":
            lines.extend([block.text, ""])
        else:
            lines.extend([block.text, ""])
    return "\n".join(lines).strip() + "\n" if lines else ""


def _render_txt(document: Document, destination: Path) -> None:
    destination.write_text(document_to_text(document), encoding="utf-8")


def _render_csv(document: Document, destination: Path) -> None:
    tables = [block.rows for block in document.blocks if block.kind == "table" and block.rows]
    with destination.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        if not tables:
            # A document with no table still has a sensible one-column spelling;
            # writing an empty file would be a worse answer than a plain list.
            for block in document.blocks:
                if block.text:
                    writer.writerow([block.text])
            return
        for index, rows in enumerate(tables):
            if index:
                writer.writerow([])
            writer.writerows(rows)


def _html_escape(text: str) -> str:
    return (
        text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace('"', "&quot;")
    )


def _render_html(document: Document, destination: Path) -> None:
    title = _html_escape(document.title or destination.stem)
    body: list[str] = []
    blocks = document.blocks
    index = 0
    while index < len(blocks):
        block = blocks[index]
        if block.kind == "list_item":
            # Consecutive items belong in one list rather than one list each.
            body.append("<ul>")
            while index < len(blocks) and blocks[index].kind == "list_item":
                body.append(f"  <li>{_html_escape(blocks[index].text)}</li>")
                index += 1
            body.append("</ul>")
            continue
        if block.kind == "table":
            body.append("<table>")
            for row_index, row in enumerate(block.rows):
                body.append("  <tr>")
                cell = "th" if row_index == 0 else "td"
                body.extend(f"    <{cell}>{_html_escape(c)}</{cell}>" for c in row)
                body.append("  </tr>")
            body.append("</table>")
        elif block.kind == "heading":
            level = min(max(block.level, 1), 6)
            body.append(f"<h{level}>{_html_escape(block.text)}</h{level}>")
        elif block.kind == "code":
            body.append(f"<pre>{_html_escape(block.text)}</pre>")
        elif block.text:
            body.append(f"<p>{_html_escape(block.text)}</p>")
        index += 1
    document_html = "\n".join(
        [
            "<!DOCTYPE html>",
            '<html lang="en">',
            "<head>",
            '<meta charset="utf-8">',
            f"<title>{title}</title>",
            "</head>",
            "<body>",
            *body,
            "</body>",
            "</html>",
            "",
        ]
    )
    destination.write_text(document_html, encoding="utf-8")


def _render_pdf(document: Document, destination: Path) -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import Paragraph, SimpleDocTemplate

    styles = getSampleStyleSheet()
    cell_style = ParagraphStyle("cell", parent=styles["BodyText"], fontSize=7, leading=9)
    width = A4[0] - 2 * inch
    story: list[Any] = []
    for block in document.blocks:
        if block.kind == "table":
            story.extend(_pdf_table(block.rows, cell_style, width))
        elif block.kind == "heading":
            name = f"Heading{min(max(block.level, 1), 4)}"
            story.append(Paragraph(_pdf_markup(block.text), styles[name]))
        elif block.kind == "list_item":
            story.append(Paragraph(f"&bull; {_pdf_markup(block.text)}", styles["BodyText"]))
        elif block.kind == "code":
            story.append(Paragraph(_pdf_markup(block.text).replace("\n", "<br/>"), styles["Code"]))
        elif block.text:
            story.append(Paragraph(_pdf_markup(block.text), styles["BodyText"]))
    if not story:
        story.append(Paragraph("(no readable content)", styles["BodyText"]))
    SimpleDocTemplate(
        str(destination),
        pagesize=A4,
        title=document.title or destination.stem,
    ).build(story)


def _pdf_markup(text: str) -> str:
    """Escape for reportlab's mini-XML, which parses the string it is given."""
    escaped = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    return escaped.strip() or "&nbsp;"


def _pdf_table(rows: list[list[str]], cell_style: Any, width: float) -> list[Any]:
    from reportlab.platypus import Paragraph, Spacer, Table, TableStyle

    columns = max(len(row) for row in rows)
    data = [
        [Paragraph(_pdf_markup(cell), cell_style) for cell in row + [""] * (columns - len(row))]
        for row in rows
    ]
    table = Table(data, colWidths=[width / columns] * columns, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("GRID", (0, 0), (-1, -1), 0.25, (0.6, 0.6, 0.6)),
                ("BACKGROUND", (0, 0), (-1, 0), (0.92, 0.92, 0.92)),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]
        )
    )
    return [table, Spacer(1, 6)]


def _render_xlsx(document: Document, destination: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    book = Workbook()
    book.remove(book.active)
    tables = 0
    heading = ""
    for block in document.blocks:
        if block.kind == "heading":
            heading = block.text
            continue
        if block.kind != "table" or not block.rows:
            continue
        tables += 1
        # Excel truncates sheet names at 31 characters and rejects a few
        # punctuation marks, so a working name is derived rather than trusted.
        name = re.sub(r"[\\/*?:\[\]]", "-", heading or f"Table {tables}")[:31] or f"Table {tables}"
        sheet = book.create_sheet(title=name)
        for row in block.rows:
            sheet.append(row)
        for index, cell in enumerate(sheet[1], start=1):
            cell.font = Font(bold=True)
            letter = get_column_letter(index)
            longest = max(
                (len(str(row[index - 1])) for row in block.rows if len(row) >= index),
                default=8,
            )
            sheet.column_dimensions[letter].width = min(max(longest + 2, 8), 60)
    if tables == 0:
        sheet = book.create_sheet(title="Content")
        for block in document.blocks:
            if block.text:
                sheet.append([block.text])
    book.save(str(destination))


_RENDERERS: dict[str, Callable[[Document, Path], None]] = {
    "txt": _render_txt,
    "csv": _render_csv,
    "html": _render_html,
    "pdf": _render_pdf,
    "xlsx": _render_xlsx,
}


def convert(source: Path, target: str, destination: Path) -> None:
    """Convert ``source`` into ``destination`` using only built-in machinery.

    Either this succeeds or it raises :class:`NativeConversionError`, so the
    caller has one failure to handle and can report it or fall back. A renderer
    that cannot write (an unwritable directory, a full disk) is folded into that
    type rather than escaping as a raw ``OSError``.
    """
    renderer = _RENDERERS.get(target.lower())
    if renderer is None:
        raise NativeConversionError(
            f"the built-in converter cannot write {target!r}; "
            f"it writes: {', '.join(sorted(TARGETS))}"
        )
    document = extract(source)
    if not document.blocks:
        raise NativeConversionError(f"no readable content was found in {source.name}")
    try:
        destination.parent.mkdir(parents=True, exist_ok=True)
        renderer(document, destination)
    except NativeConversionError:
        raise
    except Exception as exc:
        raise NativeConversionError(
            f"could not write {destination.name}: {type(exc).__name__}: {exc}"
        ) from exc
