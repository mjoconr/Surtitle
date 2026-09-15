"""The built-in document engine, which needs nothing installed.

These tests are the whole point of the module: they run on a machine with no
LibreOffice, no Office suite and no network, and prove that the formats people
actually exchange can still be read and written.
"""

from __future__ import annotations

import zipfile
from pathlib import Path

import pytest
from pypdf import PdfReader

from surtitle.tools import document_native as native

DOCX_BODY = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
 <w:body>
  <w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr>
       <w:r><w:t>Quarterly Report</w:t></w:r></w:p>
  <w:p><w:r><w:t>Revenue rose 12% &amp; costs held flat.</w:t></w:r></w:p>
  <w:tbl>
   <w:tr><w:tc><w:p><w:r><w:t>Region</w:t></w:r></w:p></w:tc>
         <w:tc><w:p><w:r><w:t>Revenue</w:t></w:r></w:p></w:tc></w:tr>
   <w:tr><w:tc><w:p><w:r><w:t>EMEA</w:t></w:r></w:p></w:tc>
         <w:tc><w:p><w:r><w:t>4.2M</w:t></w:r></w:p></w:tc></w:tr>
  </w:tbl>
  <w:p><w:pPr><w:pStyle w:val="ListParagraph"/></w:pPr>
       <w:r><w:t>Hire two engineers</w:t></w:r></w:p>
 </w:body>
</w:document>
"""

DOCX_CORE = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<cp:coreProperties
  xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties"
  xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:title>Quarterly Report</dc:title>
</cp:coreProperties>
"""

SLIDE = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<p:sld xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"
       xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
 <p:cSld><p:spTree><p:sp><p:txBody>
  <a:p><a:r><a:t>{title}</a:t></a:r></a:p>
  <a:p><a:r><a:t>{body}</a:t></a:r></a:p>
 </p:txBody></p:sp></p:spTree></p:cSld>
</p:sld>
"""

ODT_CONTENT = """<?xml version="1.0" encoding="UTF-8"?>
<office:document-content
  xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0"
  xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"
  xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0">
 <office:body><office:text>
  <text:h text:outline-level="1">Annual Review</text:h>
  <text:p>Body text here.</text:p>
  <table:table>
   <table:table-row>
    <table:table-cell><text:p>Metric</text:p></table:table-cell>
    <table:table-cell><text:p>Value</text:p></table:table-cell>
   </table:table-row>
   <table:table-row>
    <table:table-cell><text:p>Churn</text:p></table:table-cell>
    <table:table-cell><text:p>2.1%</text:p></table:table-cell>
   </table:table-row>
  </table:table>
 </office:text></office:body>
</office:document-content>
"""


def _package(path: Path, members: dict[str, str]) -> Path:
    with zipfile.ZipFile(path, "w") as bundle:
        for name, content in members.items():
            bundle.writestr(name, content)
    return path


@pytest.fixture
def docx(tmp_path) -> Path:
    return _package(
        tmp_path / "report.docx",
        {"word/document.xml": DOCX_BODY, "docProps/core.xml": DOCX_CORE},
    )


@pytest.fixture
def pptx(tmp_path) -> Path:
    return _package(
        tmp_path / "deck.pptx",
        {
            "ppt/slides/slide1.xml": SLIDE.format(title="Roadmap", body="Ship it"),
            "ppt/slides/slide2.xml": SLIDE.format(title="Risks", body="Fidelity is lower"),
        },
    )


@pytest.fixture
def odt(tmp_path) -> Path:
    return _package(tmp_path / "review.odt", {"content.xml": ODT_CONTENT})


@pytest.fixture
def xlsx(tmp_path) -> Path:
    from openpyxl import Workbook

    book = Workbook()
    first = book.active
    first.title = "Q1"
    first.append(["Region", "Revenue"])
    first.append(["EMEA", 4200])
    first.append(["APAC", 3100])
    second = book.create_sheet("Q2")
    second.append(["Region", "Revenue"])
    second.append(["EMEA", 4800])
    path = tmp_path / "book.xlsx"
    book.save(path)
    return path


def _text(path: Path) -> str:
    return native.document_to_text(native.extract(path))


class TestWordDocuments:
    def test_headings_paragraphs_tables_and_lists_all_survive(self, docx):
        kinds = [block.kind for block in native.extract(docx).blocks]
        assert kinds == ["heading", "paragraph", "table", "list_item"]

    def test_a_heading_keeps_its_level(self, docx):
        heading = native.extract(docx).blocks[0]
        assert heading.text == "Quarterly Report"
        assert heading.level == 1

    def test_entities_are_unescaped(self, docx):
        assert "Revenue rose 12% & costs held flat." in _text(docx)

    def test_a_table_keeps_its_rows_and_cells(self, docx):
        table = next(b for b in native.extract(docx).blocks if b.kind == "table")
        assert table.rows == [["Region", "Revenue"], ["EMEA", "4.2M"]]

    def test_the_title_comes_from_the_package(self, docx):
        assert native.extract(docx).title == "Quarterly Report"

    def test_a_corrupt_package_is_reported_rather_than_raised(self, tmp_path):
        broken = tmp_path / "broken.docx"
        broken.write_bytes(b"this is not a zip file")
        with pytest.raises(native.NativeConversionError):
            native.extract(broken)

    def test_an_implausibly_large_part_is_refused(self, tmp_path, monkeypatch):
        """A tiny package can claim a huge part; it must not be decompressed."""
        monkeypatch.setattr(native, "_MAX_MEMBER_BYTES", 1024)
        bomb = _package(
            tmp_path / "bomb.docx",
            {"word/document.xml": "<w:document>" + "a" * 8192 + "</w:document>"},
        )
        with pytest.raises(native.NativeConversionError, match="expands to"):
            native.extract(bomb)


class TestSpreadsheets:
    def test_csv_becomes_one_table(self, tmp_path):
        path = tmp_path / "data.csv"
        path.write_text("name,qty\nwidget,3\n", encoding="utf-8")
        table = native.extract(path).blocks[0]
        assert table.kind == "table"
        assert table.rows == [["name", "qty"], ["widget", "3"]]

    def test_each_sheet_becomes_its_own_table(self, xlsx):
        blocks = native.extract(xlsx).blocks
        assert [b.kind for b in blocks] == ["heading", "table", "heading", "table"]
        assert [b.text for b in blocks if b.kind == "heading"] == ["Q1", "Q2"]

    def test_numbers_are_stringified_not_lost(self, xlsx):
        table = native.extract(xlsx).blocks[1]
        assert table.rows[1] == ["EMEA", "4200"]


class TestPresentations:
    def test_slides_become_headings_and_paragraphs(self, pptx):
        blocks = native.extract(pptx).blocks
        assert blocks[0].kind == "heading"
        assert blocks[0].text == "Roadmap"
        assert any(b.text == "Ship it" for b in blocks)


class TestOpenDocument:
    def test_text_and_tables_are_read(self, odt):
        blocks = native.extract(odt).blocks
        assert blocks[0].kind == "heading"
        assert blocks[0].text == "Annual Review"
        table = next(b for b in blocks if b.kind == "table")
        assert table.rows == [["Metric", "Value"], ["Churn", "2.1%"]]

    def test_cells_are_not_also_emitted_as_loose_paragraphs(self, odt):
        """Descending into a table would duplicate every cell."""
        paragraphs = [b.text for b in native.extract(odt).blocks if b.kind == "paragraph"]
        assert paragraphs == ["Body text here."]


class TestMarkup:
    def test_markdown_headings_and_bullets(self, tmp_path):
        path = tmp_path / "notes.md"
        path.write_text("# Title\n\nProse.\n\n- one\n- two\n", encoding="utf-8")
        blocks = native.extract(path).blocks
        assert [b.kind for b in blocks] == ["heading", "paragraph", "list_item", "list_item"]

    def test_html_keeps_structure_and_ignores_script(self, tmp_path):
        path = tmp_path / "page.html"
        path.write_text(
            "<html><body><h1>Heading</h1><p>Para <b>bold</b>.</p>"
            "<ul><li>alpha</li><li>beta</li></ul>"
            "<table><tr><th>K</th><th>V</th></tr><tr><td>a</td><td>1</td></tr></table>"
            "<script>ignore me</script></body></html>",
            encoding="utf-8",
        )
        blocks = native.extract(path).blocks
        assert [b.kind for b in blocks] == [
            "heading",
            "paragraph",
            "list_item",
            "list_item",
            "table",
        ]
        assert "ignore me" not in _text(path)
        assert blocks[1].text == "Para bold."


class TestRendering:
    def test_pdf_is_a_real_pdf_with_the_text_in_it(self, docx, tmp_path):
        out = tmp_path / "out.pdf"
        native.convert(docx, "pdf", out)
        reader = PdfReader(str(out))
        text = "".join(page.extract_text() or "" for page in reader.pages)
        assert "Quarterly Report" in text

    def test_html_escapes_markup_rather_than_emitting_it(self, tmp_path):
        source = tmp_path / "raw.txt"
        source.write_text("a < b & c > d", encoding="utf-8")
        out = tmp_path / "out.html"
        native.convert(source, "html", out)
        body = out.read_text(encoding="utf-8")
        assert "a &lt; b &amp; c &gt; d" in body
        assert "<title>raw</title>" in body

    def test_csv_carries_the_table(self, docx, tmp_path):
        out = tmp_path / "out.csv"
        native.convert(docx, "csv", out)
        assert "EMEA,4.2M" in out.read_text(encoding="utf-8")

    def test_csv_from_a_document_with_no_table_is_not_empty(self, tmp_path):
        """An empty file would be a worse answer than a one-column list."""
        source = tmp_path / "prose.txt"
        source.write_text("First line.\n\nSecond line.\n", encoding="utf-8")
        out = tmp_path / "out.csv"
        native.convert(source, "csv", out)
        assert "First line." in out.read_text(encoding="utf-8")

    def test_xlsx_round_trips_a_table(self, docx, tmp_path):
        from openpyxl import load_workbook

        out = tmp_path / "out.xlsx"
        native.convert(docx, "xlsx", out)
        sheet = load_workbook(str(out)).worksheets[0]
        rows = [[cell.value for cell in row] for row in sheet.iter_rows()]
        assert rows[0] == ["Region", "Revenue"]
        assert rows[-1][0] == "EMEA"

    def test_txt_keeps_headings_and_marks_bullets(self, docx):
        text = _text(docx)
        assert "Quarterly Report" in text
        assert "- Hire two engineers" in text


class TestCapabilities:
    def test_office_to_pdf_is_supported(self):
        assert native.can_convert(".docx", "pdf")
        assert native.can_convert(".xlsx", "csv")

    def test_targets_libreoffice_owns_are_not_claimed(self):
        assert not native.can_convert(".docx", "rtf")
        assert not native.can_convert(".txt", "pptx")

    def test_an_unknown_source_is_rejected_with_a_reason(self, tmp_path):
        path = tmp_path / "thing.zzz"
        path.write_text("?", encoding="utf-8")
        with pytest.raises(native.NativeConversionError, match="not readable"):
            native.extract(path)

    def test_an_unsupported_target_lists_what_is_written(self, docx, tmp_path):
        with pytest.raises(native.NativeConversionError, match="cannot write"):
            native.convert(docx, "rtf", tmp_path / "out.rtf")

    def test_a_document_with_no_text_is_reported(self, tmp_path):
        empty = _package(
            tmp_path / "empty.docx", {"word/document.xml": "<w:document xmlns:w='x'/>"}
        )
        with pytest.raises(native.NativeConversionError, match="no readable content"):
            native.convert(empty, "txt", tmp_path / "out.txt")

    def test_convert_creates_missing_parent_directories(self, docx, tmp_path):
        out = tmp_path / "nested" / "deeper" / "out.txt"
        native.convert(docx, "txt", out)
        assert out.is_file()

    def test_an_unwritable_destination_is_reported_as_a_conversion_failure(self, docx, tmp_path):
        """One failure type, so callers never see a raw OSError from a renderer."""
        blocker = tmp_path / "blocker"
        blocker.write_text("I am a file, not a directory", encoding="utf-8")
        with pytest.raises(native.NativeConversionError, match="could not write"):
            native.convert(docx, "txt", blocker / "out.txt")
