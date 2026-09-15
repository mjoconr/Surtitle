"""Tests for artifact generation.

These verify the produced files are genuinely valid — the PDF is re-parsed with
pypdf and the workbook with openpyxl — because "the tool returned ok" is a much
weaker claim than "the user can open the file".
"""

from __future__ import annotations

import pytest

from surtitle.tools.artifacts import make_chart, make_pdf, make_spreadsheet
from surtitle.tools.fs_tools import ToolContext, read_file


@pytest.fixture
def ctx(tmp_path):
    return ToolContext(root=tmp_path, session_id="s1")


class TestMakePdf:
    def test_creates_a_parseable_pdf(self, ctx):
        result = make_pdf(
            ctx,
            "report.pdf",
            title="Quarterly Report",
            subtitle="Q3 2026",
            blocks=[
                {"type": "heading", "text": "Summary", "level": 1},
                {"type": "paragraph", "text": "Revenue grew eight percent."},
                {"type": "bullet_list", "items": ["Growth in EMEA", "Costs flat"]},
            ],
        )
        assert result.ok, result.error
        assert result.artifacts == ["report.pdf"]

        from pypdf import PdfReader

        reader = PdfReader(str(ctx.root / "report.pdf"))
        assert len(reader.pages) >= 1
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
        assert "Quarterly Report" in text
        assert "Revenue grew eight percent." in text
        assert "Growth in EMEA" in text

    def test_renders_a_table(self, ctx):
        result = make_pdf(
            ctx,
            "table.pdf",
            title="Data",
            blocks=[
                {
                    "type": "table",
                    "headers": ["Region", "Revenue"],
                    "rows": [["EMEA", "1.2M"], ["APAC", "0.9M"]],
                }
            ],
        )
        assert result.ok, result.error
        from pypdf import PdfReader

        text = "\n".join(
            page.extract_text() or "" for page in PdfReader(str(ctx.root / "table.pdf")).pages
        )
        assert "Region" in text
        assert "EMEA" in text
        assert "1.2M" in text

    def test_ragged_rows_are_padded_not_crashed(self, ctx):
        result = make_pdf(
            ctx,
            "ragged.pdf",
            title="Ragged",
            blocks=[{"type": "table", "headers": ["a", "b", "c"], "rows": [["1"], ["1", "2"]]}],
        )
        assert result.ok, result.error

    def test_page_size_a4(self, ctx):
        assert make_pdf(
            ctx, "a4.pdf", title="A4", blocks=[{"type": "paragraph", "text": "x"}], page_size="A4"
        ).ok

    def test_unknown_block_type_is_reported_as_a_warning(self, ctx):
        result = make_pdf(
            ctx,
            "warn.pdf",
            title="Warnings",
            blocks=[
                {"type": "paragraph", "text": "fine"},
                {"type": "nonsense", "text": "ignored"},
            ],
        )
        assert result.ok
        assert result.data["warnings"]

    def test_unicode_smart_quotes_do_not_break_the_build(self, ctx):
        result = make_pdf(
            ctx,
            "unicode.pdf",
            title="Unicode \u2014 test",
            # Curly quotes and an em dash are deliberate: reportlab's built-in
            # fonts are Latin-1 only, so this asserts the sanitiser runs.
            blocks=[
                {"type": "paragraph", "text": "It\u2019s \u201cquoted\u201d \u2014 dash\u2026"}
            ],
        )
        assert result.ok, result.error

    def test_escaped_characters_do_not_break_markup(self, ctx):
        result = make_pdf(
            ctx,
            "escaped.pdf",
            title="Escapes",
            blocks=[{"type": "paragraph", "text": "5 < 6 & 7 > 6"}],
        )
        assert result.ok, result.error

    def test_empty_blocks_is_refused(self, ctx):
        result = make_pdf(ctx, "empty.pdf", title="x", blocks=[])
        assert not result.ok

    def test_empty_title_is_refused(self, ctx):
        assert not make_pdf(
            ctx, "t.pdf", title="  ", blocks=[{"type": "paragraph", "text": "x"}]
        ).ok

    def test_extension_is_added_when_missing(self, ctx):
        result = make_pdf(ctx, "noext", title="T", blocks=[{"type": "paragraph", "text": "x"}])
        assert result.ok
        assert result.data["path"] == "noext.pdf"

    def test_escape_is_blocked(self, ctx):
        result = make_pdf(
            ctx, "../evil.pdf", title="T", blocks=[{"type": "paragraph", "text": "x"}]
        )
        assert not result.ok
        assert not (ctx.root.parent / "evil.pdf").exists()

    def test_creates_missing_directories(self, ctx):
        assert make_pdf(
            ctx, "deep/nested/r.pdf", title="T", blocks=[{"type": "paragraph", "text": "x"}]
        ).ok

    def test_image_block_embeds_a_chart(self, ctx):
        chart = make_chart(
            ctx,
            "chart.png",
            chart_type="bar",
            series=[{"name": "Revenue", "x": ["Q1", "Q2"], "y": [1, 2]}],
        )
        assert chart.ok, chart.error

        result = make_pdf(
            ctx,
            "with-chart.pdf",
            title="Chart embedded",
            blocks=[{"type": "image", "path": "chart.png", "caption": "Revenue"}],
        )
        assert result.ok, result.error
        from pypdf import PdfReader

        reader = PdfReader(str(ctx.root / "with-chart.pdf"))
        assert len(reader.pages[0].images) >= 1

    def test_missing_image_is_a_warning_not_a_failure(self, ctx):
        result = make_pdf(
            ctx,
            "noimg.pdf",
            title="T",
            blocks=[{"type": "image", "path": "nope.png"}, {"type": "paragraph", "text": "ok"}],
        )
        assert result.ok
        assert any("not found" in w for w in result.data["warnings"])


class TestMakeSpreadsheet:
    def test_creates_a_readable_workbook(self, ctx):
        result = make_spreadsheet(
            ctx,
            "data.xlsx",
            sheets=[
                {
                    "name": "Revenue",
                    "headers": ["Region", "Q1", "Q2"],
                    "rows": [["EMEA", 1.2, 1.4], ["APAC", 0.9, 1.1]],
                }
            ],
        )
        assert result.ok, result.error
        assert result.artifacts == ["data.xlsx"]

        from openpyxl import load_workbook

        workbook = load_workbook(str(ctx.root / "data.xlsx"))
        sheet = workbook["Revenue"]
        assert sheet["A1"].value == "Region"
        assert sheet["A2"].value == "EMEA"
        assert sheet["B2"].value == 1.2
        assert sheet.max_row == 3

    def test_multiple_sheets(self, ctx):
        result = make_spreadsheet(
            ctx,
            "multi.xlsx",
            sheets=[
                {"name": "One", "rows": [[1, 2]]},
                {"name": "Two", "rows": [[3, 4]]},
            ],
        )
        assert result.ok
        from openpyxl import load_workbook

        assert load_workbook(str(ctx.root / "multi.xlsx")).sheetnames == ["One", "Two"]

    def test_numeric_types_are_preserved(self, ctx):
        make_spreadsheet(ctx, "n.xlsx", sheets=[{"name": "S", "rows": [[1, 2.5, True]]}])
        from openpyxl import load_workbook

        sheet = load_workbook(str(ctx.root / "n.xlsx"))["S"]
        assert sheet["A1"].value == 1
        assert sheet["B1"].value == 2.5
        assert sheet["C1"].value is True

    def test_invalid_sheet_name_characters_are_sanitised(self, ctx):
        result = make_spreadsheet(ctx, "bad.xlsx", sheets=[{"name": "a/b:c*d?e[f]", "rows": [[1]]}])
        assert result.ok
        from openpyxl import load_workbook

        name = load_workbook(str(ctx.root / "bad.xlsx")).sheetnames[0]
        assert not any(ch in name for ch in "[]:*?/\\")

    def test_unserialisable_values_are_stringified(self, ctx):
        result = make_spreadsheet(
            ctx, "obj.xlsx", sheets=[{"name": "S", "rows": [[{"a": 1}, [1, 2]]]}]
        )
        assert result.ok, result.error

    def test_empty_sheets_is_refused(self, ctx):
        assert not make_spreadsheet(ctx, "e.xlsx", sheets=[]).ok

    def test_extension_is_added(self, ctx):
        result = make_spreadsheet(ctx, "noext", sheets=[{"name": "S", "rows": [[1]]}])
        assert result.ok
        assert result.data["path"] == "noext.xlsx"

    def test_escape_is_blocked(self, ctx):
        assert not make_spreadsheet(ctx, "../evil.xlsx", sheets=[{"name": "S", "rows": [[1]]}]).ok

    def test_round_trips_through_read_file(self, ctx):
        """A spreadsheet this tool wrote can be read straight back."""
        make_spreadsheet(ctx, "rt.xlsx", sheets=[{"name": "S", "headers": ["a"], "rows": [[1]]}])
        result = read_file(ctx, "rt.xlsx")
        assert result.ok, result.error
        assert result.data["kind"] == "office"
        assert "a" in result.data["content"]


class TestMakeChart:
    def test_creates_a_png(self, ctx):
        result = make_chart(
            ctx,
            "chart.png",
            chart_type="line",
            title="Trend",
            series=[{"name": "Revenue", "x": [1, 2, 3], "y": [10, 20, 30]}],
        )
        assert result.ok, result.error
        path = ctx.root / "chart.png"
        assert path.exists()
        assert path.read_bytes()[:8] == b"\x89PNG\r\n\x1a\n"

    @pytest.mark.parametrize("kind", ["line", "bar", "hbar", "pie", "scatter", "area", "hist"])
    def test_every_chart_type_renders(self, ctx, kind):
        result = make_chart(
            ctx,
            f"{kind}.png",
            chart_type=kind,
            series=[{"name": "S", "x": ["A", "B", "C"], "y": [3, 1, 2]}],
        )
        assert result.ok, result.error

    def test_category_labels_without_x_values(self, ctx):
        assert make_chart(ctx, "s.png", chart_type="bar", series=[{"name": "S", "y": [1, 2, 3]}]).ok

    def test_unsupported_type_is_refused_with_options(self, ctx):
        result = make_chart(ctx, "x.png", chart_type="donut", series=[{"y": [1]}])
        assert not result.ok
        assert "pie" in result.error

    def test_empty_series_is_refused(self, ctx):
        assert not make_chart(ctx, "x.png", chart_type="line", series=[]).ok

    def test_series_without_y_is_skipped(self, ctx):
        result = make_chart(ctx, "noy.png", chart_type="line", series=[{"name": "empty"}])
        # No data plotted, but the figure is still valid; must not raise.
        assert isinstance(result.ok, bool)

    def test_non_numeric_values_report_an_error(self, ctx):
        result = make_chart(ctx, "bad.png", chart_type="line", series=[{"y": ["not", "numbers"]}])
        assert not result.ok

    def test_svg_output(self, ctx):
        result = make_chart(ctx, "chart.svg", chart_type="bar", series=[{"name": "S", "y": [1, 2]}])
        assert result.ok, result.error
        assert b"<svg" in (ctx.root / "chart.svg").read_bytes()[:400]

    def test_escape_is_blocked(self, ctx):
        assert not make_chart(ctx, "../evil.png", chart_type="line", series=[{"y": [1]}]).ok
